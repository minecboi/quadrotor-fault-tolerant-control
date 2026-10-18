#!/usr/bin/env python3
"""
scripts/update_future_work.py
Updates the Future Work / Future Scope sections in REPORT.md, README.md, and
Quadrotor Single Motor Failure Report.docx, and re-exports the PDF via LibreOffice.
Adds:
1. Power loss estimation improvements (tilt-based observer feedback suppression and multi-observer fusion).
2. Completing and hardening the active recovery system (resolving transient handover tumbles via native onboard integration).
"""
import docx
import os
import shutil
import subprocess
import re

BASE_DIR = "/home/luffy/chayan_2502MT04_mini_ps_ideaforge -20261016T131509Z-1-001/chayan_2502MT04_mini_ps_ideaforge "
DOCX_PATH = os.path.join(BASE_DIR, "Quadrotor Single Motor Failure Report.docx")
MD_PATH = os.path.join(BASE_DIR, "REPORT.md")
README_PATH = os.path.join(BASE_DIR, "README.md")

def update_markdown():
    # 1. Update REPORT.md
    with open(MD_PATH, 'r') as f:
        text = f.read()

    old_future_md = r"### What is Left / Recommended for Future Development:.*?(?=\Z)"
    new_future_md = """### What is Left / Recommended for Future Development:
1. **Advanced Power Loss Estimation via Multi-Observer Fusion:**
   Currently, the rapid tilt-based kinematic estimator ($E_1$) suffers from closed-loop attitude-hold attenuation (PX4's inner rate loop aggressively commands counter-torques to keep the vehicle level, suppressing tilt $\\Delta \\theta \\approx 0^\\circ$ and underestimating severity during the initial $350\\,\\text{ms}$ window), while the opposite-motor command ratio estimator ($E_2$) requires $1.0 - 2.0\\,\\text{s}$ for mixer integrators to settle. Future work will develop an augmented Extended Kalman Filter (EKF) or adaptive aerodynamic disturbance observer that explicitly incorporates inner-loop torque setpoint integration and control allocation residuals, enabling accurate, fast-converging thrust loss estimation ($\\hat{\\rho}$) within $100 - 200\\,\\text{ms}$ without relying solely on passive unconstrained rigid-body tilt.
2. **Completing and Hardening the Closed-Loop Active Recovery System:**
   While the fault detection and 4-quadrant isolation engine achieved 100% precision across live SITL runs, the active offboard recovery controller and supervisory handover pipeline experienced transient tumble/altitude drop failures during specific high-speed dynamic maneuvers. Completing the recovery system entails:
   - Eliminating MAVLink offboard message latency by implementing native onboard PX4 mixer/allocation handover directly inside the flight controller firmware.
   - Implementing bumpless transfer and feedforward attitude-state initialization for the spinning relaxed hover manifold during the 1.2-second grace period.
   - Tuning actuator anti-windup to prevent intermediate tumble when transitioning between ADAPTIVE level control and YAW-FREE spinning mode.
3. **Continuous Manifold Trajectory Blending:** Implement continuous Lyapunov-based manifold transitions across the 30%–40% boundary to smoothly transition between level hover and spinning flight.
4. **Onboard PX4 Native C++ uORB Module:** Compile the 2-step residual observer directly into a PX4 C++ module (`src/modules/fault_detector`) for microsecond execution directly on flight controller hardware.
5. **Physical Hardware Flight Test:** Execute the proposed Track B tethered testing protocol on physical quadrotor hardware.
"""
    text = re.sub(old_future_md, lambda _: new_future_md, text, flags=re.DOTALL)
    text = text.replace("5. **Full Live Telemetry Dataset (30 Runs):**", "5. **Full Live Telemetry Dataset (22 Evaluated Runs):**")

    with open(MD_PATH, 'w') as f:
        f.write(text)
    print("Updated REPORT.md future scope successfully!")

    # 2. Update README.md
    with open(README_PATH, 'r') as f:
        readme = f.read()

    old_future_readme = r"### 3\. Future Scope / Next Steps.*?(?=---|\Z)"
    new_future_readme = """### 3. Future Scope / Next Steps
- **Advanced Power Loss Estimation via Multi-Observer Fusion:** Advance real-time fault severity estimation by fusing the fast tilt-kinematic observer with inner-loop control torque residuals and EKF state estimation. This overcomes closed-loop attitude-hold tilt suppression and eliminates the 1–2s integrator settling delay of the opposite-motor estimator.
- **Completing and Hardening the Closed-Loop Recovery System:** Complete and stabilize the active recovery pipeline (which experienced transient tumbles during live offboard handover) by replacing offboard MAVLink overrides with native PX4 onboard allocation switching, implementing bumpless state handover during the 1.2s grace period, and tuning anti-windup.
- **Continuous Manifold Trajectory Blending:** Implement continuous Lyapunov-based manifold transitions across the 30%–40% boundary to smoothly transition between level hover and spinning flight.
- **Embedded C++ Implementation:** Port the residual observer into a native PX4 uORB C++ module for sub-millisecond execution directly on the flight controller.
- **Physical Hardware Tethered Flight:** Test the Track B hardware build using the proposed dual-swivel tether safety rig.

"""
    readme = re.sub(old_future_readme, lambda _: new_future_readme, readme, flags=re.DOTALL)
    with open(README_PATH, 'w') as f:
        f.write(readme)
    print("Updated README.md future scope successfully!")

def update_docx():
    doc = docx.Document(DOCX_PATH)

    # Locate Bonus 4 paragraph
    for i, p in enumerate(doc.paragraphs):
        if "Bonus 4: Future Work on the Handover Boundary" in p.text:
            next_p = doc.paragraphs[i+1]
            next_p.text = (
                "Based on our live flight testing and the empirical discovery of the 30% to 40% effectiveness boundary, we propose five key directions for future development:\n"
                "1. Advanced Power Loss Estimation via Multi-Observer Fusion: Currently, the rapid tilt-based kinematic estimator (E1) suffers from closed-loop feedback attenuation (PX4's inner rate loop aggressively keeps the vehicle level, suppressing tilt Δθ to zero during the initial 350 ms window), while the opposite-motor command ratio (E2) requires 1.0–2.0 s for mixer integrators to settle. Future work will develop an augmented Extended Kalman Filter (EKF) combining high-rate torque setpoint residuals and actuator commands, achieving converged thrust loss estimation within 100–200 ms.\n"
                "2. Completing and Hardening the Closed-Loop Recovery System: While our live detection and 4-quadrant isolation achieved 100% accuracy, the active recovery handover controller experienced transient tumble and altitude loss failures during aggressive dynamic flight transitions. Future work will complete this pipeline by compiling the supervisor into native onboard PX4 firmware (eliminating offboard MAVLink latency), adding feedforward attitude-state initialization for the relaxed-spin manifold, and tuning anti-windup to guarantee bumpless controller transfer.\n"
                "3. Smooth Trajectory Blending Across the 30%–40% Zone: Currently, the supervisory switch between level flight and spinning hover is discrete. Future work should investigate continuous Lyapunov-based manifold transitions so that if a motor progressively degrades from 50% down to 20%, the aircraft transitions into a spin without transient altitude loss.\n"
                "4. Crosswind & Turbulence Robustness: The 40% margin was identified in clean, deterministic simulation. Evaluating how ambient wind gusts shift this boundary (e.g. whether a 10 m/s wind pushes the threshold up to 50%) is a key step before hardware deployment.\n"
                "5. Flight Test Validation on Physical Rig: Validating the opposite-motor command ratio and the 40% level-flight boundary on a physical quadrotor with throttle-capped ESCs under tethered safety."
            )
            print("Updated Bonus 4 in DOCX successfully!")
            break

    doc.save(DOCX_PATH)
    shutil.copyfile(DOCX_PATH, os.path.join(BASE_DIR, "docs/Quadrotor Single Motor Failure Report.docx"))

def export_pdf():
    cmd = 'libreoffice --headless --convert-to pdf "Quadrotor Single Motor Failure Report.docx"'
    res = subprocess.run(cmd, shell=True, cwd=BASE_DIR, capture_output=True, text=True)
    if res.returncode == 0:
        shutil.copyfile(
            os.path.join(BASE_DIR, "Quadrotor Single Motor Failure Report.pdf"),
            os.path.join(BASE_DIR, "Chayan_2502MT04_Quadrotor Single Motor Failure Report.pdf")
        )
        print("Exported PDF successfully!")
    else:
        print("PDF export failed:", res.stderr)

if __name__ == "__main__":
    update_markdown()
    update_docx()
    export_pdf()
